from docx import Document
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# Create a new Document
doc = Document()

# Create a table
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Header 1"
table.cell(0, 1).text = "Header 2"
table.cell(1, 0).text = "Row 1, Cell 1"
table.cell(1, 1).text = "Row 1, Cell 2"

# Set background color for a cell
cell = table.cell(1, 1)
cell._element.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="FFFF00"/>'.format(nsdecls('w'))))  # Yellow background

# Save the document
doc.save("output/06_Formatting_tables.docx")
