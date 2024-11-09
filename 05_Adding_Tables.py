from docx import Document

# Create a new Document
doc = Document()

# Create a table
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Header 1"
table.cell(0, 1).text = "Header 2"
table.cell(1, 0).text = "Row 1, Cell 1"
table.cell(1, 1).text = "Row 1, Cell 2"

# Save the document
doc.save("output/05_Adding_Tables.docx")
